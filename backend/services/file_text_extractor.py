# backend/services/file_text_extractor.py

import os
from pathlib import Path
import tempfile

import fitz  # PyMuPDF
from docx import Document
from pptx import Presentation


def get_extension(filename: str) -> str:
    return Path(filename).suffix.lower()


def extract_text_from_pdf(file_bytes: bytes) -> str:
    text_parts = []
    with fitz.open(stream=file_bytes, filetype="pdf") as doc:
        for page in doc:
            page_text = page.get_text("text")
            if page_text:
                text_parts.append(page_text.strip())
    return "\n\n".join(part for part in text_parts if part)


def extract_text_from_docx(file_bytes: bytes) -> str:
    with tempfile.NamedTemporaryFile(delete=False, suffix=".docx") as tmp:
        tmp.write(file_bytes)
        tmp_path = tmp.name

    try:
        doc = Document(tmp_path)
        paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        return "\n".join(paragraphs)
    finally:
        if os.path.exists(tmp_path):
            os.remove(tmp_path)


def extract_text_from_pptx(file_bytes: bytes) -> str:
    with tempfile.NamedTemporaryFile(delete=False, suffix=".pptx") as tmp:
        tmp.write(file_bytes)
        tmp_path = tmp.name

    try:
        prs = Presentation(tmp_path)
        text_parts = []

        for slide in prs.slides:
            slide_parts = []
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text:
                    value = shape.text.strip()
                    if value:
                        slide_parts.append(value)
            if slide_parts:
                text_parts.append("\n".join(slide_parts))

        return "\n\n".join(text_parts)
    finally:
        if os.path.exists(tmp_path):
            os.remove(tmp_path)


def extract_text_by_filename(filename: str, file_bytes: bytes) -> str:
    extension = get_extension(filename)

    if extension == ".pdf":
        return extract_text_from_pdf(file_bytes)
    if extension == ".docx":
        return extract_text_from_docx(file_bytes)
    if extension == ".pptx":
        return extract_text_from_pptx(file_bytes)

    raise ValueError("Unsupported file type for extraction.")